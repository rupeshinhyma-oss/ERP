"""
Monthly Payroll Excel Export Engine for Inhyma ERP HRMS.

Builds an authoritative, audit-ready, high-aesthetic Excel (.xlsx) workbook directly
from PostgreSQL payroll records and calculated items without recalculating or modifying payroll data.
Supports:
- All authorized employees in the selected payroll month
- Single employee filtered export
- Dynamic configured earnings and deductions (Basic, HRA, PF, PT, Special Allowance, etc.)
- Granular RBAC enforcement (safeguarding unapproved drafts and unauthorized salary access)
"""

from __future__ import annotations

import io
import uuid
from datetime import date, datetime, timedelta, timezone
from typing import Optional, Set

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload
from fastapi import HTTPException, status

from app.hrms.models import (
    HrmsEmployeeLocation,
    HrmsLocation,
    HrmsMonthlyPayroll,
    HrmsMonthlyPayrollItem,
    HrmsPayrollComponent,
)
from app.rbac.models import Role, UserRole
from app.users.models import User

# Indian Standard Time (Asia/Kolkata)
try:
    from zoneinfo import ZoneInfo
    IST = ZoneInfo("Asia/Kolkata")
except Exception:
    IST = timezone(timedelta(hours=5, minutes=30))


def to_ist(dt: datetime | None) -> datetime | None:
    if not dt:
        return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(IST)


def format_dt_ist(dt: datetime | None) -> str:
    ist_dt = to_ist(dt)
    return ist_dt.strftime("%d-%b-%Y %I:%M %p") if ist_dt else "-"


async def build_monthly_payroll_excel(
    db: AsyncSession,
    payroll_month: str,
    employee_id: Optional[uuid.UUID] = None,
    allowed_employee_ids: Optional[Set[uuid.UUID]] = None,
    is_admin: bool = False,
    require_approved_only: bool = False,
) -> tuple[io.BytesIO, int, str]:
    """
    Generate authoritative payroll Excel workbook directly from PostgreSQL payroll records.
    Returns: (BytesIO buffer, item_count, filename).
    """
    cleaned_month = payroll_month.strip()

    # 1. Fetch monthly payroll snapshot
    query = (
        select(HrmsMonthlyPayroll)
        .options(
            selectinload(HrmsMonthlyPayroll.items).selectinload(HrmsMonthlyPayrollItem.employee),
            selectinload(HrmsMonthlyPayroll.items).selectinload(HrmsMonthlyPayrollItem.salary),
        )
        .where(HrmsMonthlyPayroll.payroll_month == cleaned_month)
    )
    payroll = (await db.execute(query)).scalar_one_or_none()

    if not payroll or not payroll.items:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"No generated payroll records found for {cleaned_month}.",
        )

    # 2. Enforce approval safeguard for non-admin/non-HR
    if require_approved_only and payroll.status != "APPROVED":
        # Check if individual items have APPROVED status
        approved_items = [it for it in payroll.items if it.status == "APPROVED"]
        if not approved_items:
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN,
                detail=f"Payroll for {cleaned_month} is currently in {payroll.status} status and not yet approved for release.",
            )

    # 3. Filter items by employee_id or allowed_employee_ids
    items: list[HrmsMonthlyPayrollItem] = list(payroll.items)

    if employee_id:
        items = [it for it in items if it.employee_id == employee_id]
        if not items:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"No payroll record found for employee in {cleaned_month}.",
            )
    elif allowed_employee_ids is not None:
        items = [it for it in items if it.employee_id in allowed_employee_ids]

    if require_approved_only:
        items = [it for it in items if it.status == "APPROVED"]
        if not items:
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN,
                detail=f"No approved payroll records available for download in {cleaned_month}.",
            )

    # Sort items by employee code/name
    items.sort(
        key=lambda it: (
            (getattr(it.employee, "employee_code", None) or ""),
            (getattr(it.employee, "first_name", None) or ""),
        )
    )

    uids = [it.employee_id for it in items]

    # 4. Batch resolve branches and departments
    branch_map: dict[uuid.UUID, str] = {}
    dept_map: dict[uuid.UUID, str] = {}

    if uids:
        try:
            loc_stmt = (
                select(HrmsEmployeeLocation.user_id, HrmsLocation.name)
                .join(HrmsLocation, HrmsLocation.id == HrmsEmployeeLocation.location_id)
                .where(HrmsEmployeeLocation.user_id.in_(uids), HrmsLocation.deleted_at.is_(None))
                .order_by(HrmsEmployeeLocation.is_primary.desc())
            )
            for uid, lname in (await db.execute(loc_stmt)).all():
                if uid not in branch_map:
                    branch_map[uid] = lname
        except Exception:
            pass

        try:
            r_stmt = (
                select(UserRole.user_id, Role.name)
                .join(Role, Role.id == UserRole.role_id)
                .where(UserRole.user_id.in_(uids), Role.deleted_at.is_(None))
            )
            for uid, rname in (await db.execute(r_stmt)).all():
                if rname.lower() not in ("user", "employee", "department_manager", "super_admin", "admin") and uid not in dept_map:
                    dept_map[uid] = rname
        except Exception:
            pass

    # 5. Approver name resolution
    approver_name = "System"
    if payroll.approved_by:
        approver = await db.get(User, payroll.approved_by)
        if approver:
            approver_name = f"{approver.first_name or ''} {approver.last_name or ''}".strip() or approver.username or "Admin"

    # 6. Dynamically collect all configured earning and deduction components present across items
    # Maintain canonical display order from HrmsPayrollComponent
    comps_stmt = select(HrmsPayrollComponent).where(HrmsPayrollComponent.deleted_at.is_(None)).order_by(HrmsPayrollComponent.display_order.asc())
    all_known_comps = list((await db.execute(comps_stmt)).scalars().all())
    comp_order = {c.code.upper(): (c.display_order, c.name) for c in all_known_comps}

    earnings_keys: dict[str, str] = {}  # code -> display name
    deductions_keys: dict[str, str] = {}  # code -> display name

    for it in items:
        if it.earnings_breakdown:
            for e in it.earnings_breakdown:
                code = (e.get("code") or e.get("name") or "OTHER").upper()
                name = e.get("name") or comp_order.get(code, (99, code))[1]
                earnings_keys[code] = name

        if it.deductions_breakdown:
            for d in it.deductions_breakdown:
                code = (d.get("code") or d.get("name") or "OTHER").upper()
                name = d.get("name") or comp_order.get(code, (99, code))[1]
                deductions_keys[code] = name

    # Sort earning and deduction keys by display_order if known
    sorted_earnings = sorted(earnings_keys.items(), key=lambda x: comp_order.get(x[0], (99, x[1]))[0])
    sorted_deductions = sorted(deductions_keys.items(), key=lambda x: comp_order.get(x[0], (99, x[1]))[0])

    # 7. Construct Workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Payroll Statement"
    ws.views.sheetView[0].showGridLines = True

    # Palette
    navy_dark = "1E293B"
    navy_fill = PatternFill(start_color=navy_dark, end_color=navy_dark, fill_type="solid")
    accent_blue = "2563EB"
    banner_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    alt_row_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    total_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

    border_side = Side(border_style="thin", color="CBD5E1")
    thin_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
    double_bottom = Border(left=border_side, right=border_side, top=border_side, bottom=Side(border_style="double", color="1E293B"))

    font_banner_title = Font(name="Calibri", size=14, bold=True, color="0F172A")
    font_banner_sub = Font(name="Calibri", size=9.5, color="475569")
    font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=9.5, bold=True, color="0F172A")
    font_regular = Font(name="Calibri", size=9.5, color="1E293B")
    font_total = Font(name="Calibri", size=10, bold=True, color="0F172A")

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # Banner Block (Rows 1 to 3)
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 12

    try:
        y, m_num = map(int, cleaned_month.split("-"))
        month_label = date(y, m_num, 1).strftime("%B %Y")
    except Exception:
        month_label = cleaned_month

    ws.cell(row=1, column=1, value=f"INHYMA ERP — MONTHLY PAYROLL STATEMENT ({month_label.upper()})").font = font_banner_title
    status_label = f"Status: {payroll.status}"
    if payroll.approved_at:
        status_label += f" | Approved by: {approver_name} on {format_dt_ist(payroll.approved_at)}"
    else:
        status_label += f" | Generated on: {format_dt_ist(payroll.processed_at)}"
    status_label += f" | Working Days: {payroll.working_days} | Total Employees: {len(items)}"
    ws.cell(row=2, column=1, value=status_label).font = font_banner_sub

    # Define Columns
    # Static info columns
    col_defs: list[tuple[str, str, str]] = [  # (header, alignment, format)
        ("Employee ID", "center", "@"),
        ("Employee Name", "left", "@"),
        ("Department", "left", "@"),
        ("Branch / Office", "left", "@"),
        ("Payroll Month", "center", "@"),
        ("Annual CTC", "right", "#,##0.00"),
        ("Monthly CTC", "right", "#,##0.00"),
        ("Working Days", "center", "0"),
        ("Present Days", "center", "0.0"),
        ("Paid Leave Days", "center", "0.0"),
        ("Holiday Days", "center", "0"),
        ("Weekend Days", "center", "0"),
        ("Paid Days", "center", "0.0"),
        ("LOP Days", "center", "0.0"),
    ]

    # Dynamic Earning columns
    for code, name in sorted_earnings:
        col_defs.append((f"{name} ({code})", "right", "#,##0.00"))

    col_defs.append(("Gross Earnings", "right", "#,##0.00"))
    col_defs.append(("LOP Deduction", "right", "#,##0.00"))

    # Dynamic Deduction columns
    for code, name in sorted_deductions:
        col_defs.append((f"{name} ({code})", "right", "#,##0.00"))

    col_defs.append(("Other Adjustments", "right", "#,##0.00"))
    col_defs.append(("Total Deductions", "right", "#,##0.00"))
    col_defs.append(("Net Payable Salary", "right", "#,##0.00"))
    col_defs.append(("Payroll Status", "center", "@"))
    col_defs.append(("Approval Info", "left", "@"))

    num_cols = len(col_defs)

    # Header Row (Row 4)
    header_row = 4
    ws.row_dimensions[header_row].height = 28
    for col_idx, (col_name, col_align, _) in enumerate(col_defs, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=col_name)
        cell.font = font_header
        cell.fill = navy_fill
        cell.alignment = align_center
        cell.border = thin_border

    # Data Rows (Row 5 onwards)
    data_start_row = 5
    for row_offset, it in enumerate(items):
        row_num = data_start_row + row_offset
        ws.row_dimensions[row_num].height = 20
        row_fill = alt_row_fill if (row_offset % 2 == 1) else PatternFill(fill_type=None)

        emp = it.employee
        emp_code = getattr(emp, "employee_code", None) or f"EMP-{str(it.employee_id)[:6].upper()}"
        first = getattr(emp, "first_name", "") or ""
        last = getattr(emp, "last_name", "") or ""
        emp_name = f"{first} {last}".strip() or getattr(emp, "display_name", None) or getattr(emp, "username", "Employee")
        dept = dept_map.get(it.employee_id) or getattr(emp, "department", "General") or "General"
        branch = branch_map.get(it.employee_id) or "Thane"

        paid_days = round(it.present_days + it.paid_leave_days + it.holiday_days + it.weekend_days, 1)

        # Map earnings
        emp_earnings: dict[str, float] = {}
        if it.earnings_breakdown:
            for e in it.earnings_breakdown:
                c_code = (e.get("code") or e.get("name") or "").upper()
                emp_earnings[c_code] = float(e.get("monthly_amount") or 0.0)

        # Map deductions
        emp_deductions: dict[str, float] = {}
        if it.deductions_breakdown:
            for d in it.deductions_breakdown:
                c_code = (d.get("code") or d.get("name") or "").upper()
                emp_deductions[c_code] = float(d.get("monthly_amount") or 0.0)

        # Additions / Adjustments
        other_adjustments = 0.0
        if it.additions_breakdown:
            for a in it.additions_breakdown:
                other_adjustments += float(a.get("amount") or 0.0)

        # Approval string
        approval_str = (
            f"Approved by {approver_name} ({to_ist(payroll.approved_at).strftime('%d-%b-%Y')})"
            if payroll.approved_at
            else "Pending Review"
        )

        row_values: list[tuple[any, str, str]] = [
            (emp_code, "center", "@"),
            (emp_name, "left", "@"),
            (dept, "left", "@"),
            (branch, "left", "@"),
            (cleaned_month, "center", "@"),
            (it.annual_ctc, "right", "#,##0.00"),
            (it.monthly_ctc, "right", "#,##0.00"),
            (it.working_days, "center", "0"),
            (it.present_days, "center", "0.0"),
            (it.paid_leave_days, "center", "0.0"),
            (it.holiday_days, "center", "0"),
            (it.weekend_days, "center", "0"),
            (paid_days, "center", "0.0"),
            (it.lop_days, "center", "0.0"),
        ]

        for code, _ in sorted_earnings:
            val = emp_earnings.get(code, 0.0)
            row_values.append((val, "right", "#,##0.00"))

        row_values.append((it.gross_amount, "right", "#,##0.00"))
        row_values.append((it.lop_deduction, "right", "#,##0.00"))

        for code, _ in sorted_deductions:
            val = emp_deductions.get(code, 0.0)
            row_values.append((val, "right", "#,##0.00"))

        row_values.append((other_adjustments, "right", "#,##0.00"))
        row_values.append((it.total_deductions, "right", "#,##0.00"))
        row_values.append((it.net_salary, "right", "#,##0.00"))
        row_values.append((it.status or payroll.status, "center", "@"))
        row_values.append((approval_str, "left", "@"))

        for col_idx, (val, col_align, num_fmt) in enumerate(row_values, start=1):
            cell = ws.cell(row=row_num, column=col_idx, value=val)
            cell.font = font_regular
            cell.border = thin_border
            if row_fill.fill_type:
                cell.fill = row_fill
            cell.alignment = align_left if col_align == "left" else (align_right if col_align == "right" else align_center)
            cell.number_format = num_fmt

    # Summary Totals Row
    data_end_row = data_start_row + len(items) - 1
    total_row = data_end_row + 1
    ws.row_dimensions[total_row].height = 24

    ws.cell(row=total_row, column=1, value="TOTALS").font = font_total
    ws.cell(row=total_row, column=1).alignment = align_center
    ws.cell(row=total_row, column=1).fill = total_fill
    ws.cell(row=total_row, column=1).border = double_bottom

    for col_idx in range(2, num_cols + 1):
        cell = ws.cell(row=total_row, column=col_idx)
        cell.fill = total_fill
        cell.border = double_bottom
        cell.font = font_total

        _, col_align, num_fmt = col_defs[col_idx - 1]
        col_letter = get_column_letter(col_idx)

        # Apply SUM formula for currency and day count columns
        if num_fmt in ("#,##0.00", "0.0", "0") and col_idx >= 6 and col_idx < num_cols - 1:
            cell.value = f"=SUM({col_letter}{data_start_row}:{col_letter}{data_end_row})"
            cell.number_format = num_fmt
            cell.alignment = align_right if col_align == "right" else align_center
        else:
            cell.value = ""

    # Freeze panes at D5 (sticky columns: Employee ID, Name, Department)
    ws.freeze_panes = "D5"

    # Auto-adjust column widths
    for col_idx, (col_name, _, _) in enumerate(col_defs, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = max(len(col_name), 10)
        # sample first 20 data rows for length
        for r in range(data_start_row, min(data_start_row + 20, data_end_row + 1)):
            v = ws.cell(row=r, column=col_idx).value
            if v is not None:
                max_len = max(max_len, len(str(v)))
        ws.column_dimensions[col_letter].width = min(max_len + 4, 35)

    # Determine filename
    target_label = "All_Employees"
    if employee_id and len(items) == 1:
        single_emp = items[0].employee
        code_part = getattr(single_emp, "employee_code", None) or f"EMP-{str(employee_id)[:6].upper()}"
        name_part = (getattr(single_emp, "first_name", "") or "Employee").replace(" ", "_")
        target_label = f"{code_part}_{name_part}"

    filename = f"Payroll_Export_{target_label}_{cleaned_month.replace('-', '_')}.xlsx"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf, len(items), filename
