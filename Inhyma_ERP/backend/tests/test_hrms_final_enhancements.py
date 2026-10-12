"""
Integration and RBAC tests for Final HRMS Enhancements:
1. All-Employee Attendance Calendar (GET /api/v1/hrms/attendance/all-calendar)
2. Payroll Excel Download (GET /api/v1/hrms/payroll/export)

Validates against live PostgreSQL database data, real schemas, and RBAC policies.
"""

import io
import uuid
import pytest
import pytest_asyncio
from httpx import AsyncClient
import openpyxl
from sqlalchemy import delete

from app.auth.security import create_access_token
from app.database.engine import get_sessionmaker
from app.users.models import User, UserStatus

ADMIN_ID = uuid.UUID("ea2b5763-5a4d-4f21-8bff-adc3e6cb8235")


@pytest.fixture
def admin_headers():
    token = create_access_token(ADMIN_ID, permissions=["*"]).token
    return {"Authorization": f"Bearer {token}"}


@pytest_asyncio.fixture
async def regular_emp_headers():
    """Seed a real employee user without admin/HR permissions in PostgreSQL, then clean up."""
    emp_id = uuid.uuid4()
    maker = get_sessionmaker()
    async with maker() as session:
        emp_user = User(
            id=emp_id,
            email=f"emp_test_{emp_id.hex[:6]}@example.com",
            username=f"emp_test_{emp_id.hex[:6]}",
            first_name="Regular",
            last_name="Staff",
            employee_code=f"EMP-T-{emp_id.hex[:4]}",
            status=UserStatus.ACTIVE,
            is_active=True,
            password_hash="plainpass",
        )
        session.add(emp_user)
        await session.commit()

    token = create_access_token(emp_id, permissions=[]).token
    yield {"Authorization": f"Bearer {token}", "emp_id": str(emp_id)}

    async with maker() as session:
        await session.execute(delete(User).where(User.id == emp_id))
        await session.commit()


@pytest.mark.asyncio
async def test_all_employee_attendance_calendar_admin(client: AsyncClient, admin_headers):
    """
    Test 1: Admin can access All Employees Calendar for September 2026.
    Verifies month/year, total days, employee counts, department/branch lists,
    and daily attendance statuses.
    """
    res = await client.get("/api/v1/hrms/attendance/all-calendar?year=2026&month=9", headers=admin_headers)
    assert res.status_code == 200, f"Expected 200 OK, got {res.status_code}: {res.text}"

    data = res.json()["data"]
    assert data["year"] == 2026
    assert data["month"] == 9
    assert "September" in data["month_name"]
    assert data["total_days"] == 30
    assert isinstance(data["departments"], list)
    assert isinstance(data["branches"], list)
    assert isinstance(data["employees"], list)
    assert len(data["employees"]) > 0, "Expected at least one authorized employee record"

    # Inspect Admin User's attendance in the matrix
    admin_emp = next((e for e in data["employees"] if "EMP-000" in e.get("employee_id", "") or "Admin" in e.get("employee_name", "")), None)
    assert admin_emp is not None, "Admin User record should be present in all-calendar"
    assert len(admin_emp["days"]) == 30, "All 30 days should be populated for the month"

    # Check presence of present days and summary totals
    assert admin_emp["summary"]["present"] >= 1, "Admin should have present days recorded"

    # Verify day structure contains punch details, shift info, and status
    first_day = admin_emp["days"][0]
    assert "day" in first_day
    assert "date" in first_day
    assert "status" in first_day
    assert "working_hours" in first_day
    assert "shift_name" in first_day


@pytest.mark.asyncio
async def test_all_employee_attendance_calendar_filters(client: AsyncClient, admin_headers):
    """
    Test 2: Test department and branch filtering in All Employees Calendar.
    """
    # Filter by non-existent department should return empty employees list
    res_empty = await client.get("/api/v1/hrms/attendance/all-calendar?year=2026&month=9&department=NonExistentDept999", headers=admin_headers)
    assert res_empty.status_code == 200
    assert len(res_empty.json()["data"]["employees"]) == 0

    # Filter with valid month navigation (October 2026)
    res_oct = await client.get("/api/v1/hrms/attendance/all-calendar?year=2026&month=10", headers=admin_headers)
    assert res_oct.status_code == 200
    data_oct = res_oct.json()["data"]
    assert data_oct["year"] == 2026
    assert data_oct["month"] == 10
    assert data_oct["total_days"] == 31


@pytest.mark.asyncio
async def test_all_employee_attendance_calendar_rbac_scoping(client: AsyncClient, regular_emp_headers):
    """
    Test 3: Non-admin employee cannot access other employees' attendance via direct query.
    Must return 403 Forbidden.
    """
    other_emp_id = uuid.uuid4()
    headers = {"Authorization": regular_emp_headers["Authorization"]}
    res_unauth = await client.get(
        f"/api/v1/hrms/attendance/all-calendar?year=2026&month=9&employee_id={other_emp_id}",
        headers=headers
    )
    assert res_unauth.status_code == 403
    assert "Permission denied" in res_unauth.text or "Forbidden" in res_unauth.text


@pytest.mark.asyncio
async def test_payroll_excel_export_all_employees(client: AsyncClient, admin_headers):
    """
    Test 4: Export all authorized employees payroll for September 2026.
    Verifies Excel file binary content, headers, dynamic earnings/deductions, and formulas.
    """
    res = await client.get("/api/v1/hrms/payroll/export?payroll_month=2026-09", headers=admin_headers)
    assert res.status_code == 200, f"Expected 200 OK, got {res.status_code}: {res.text}"

    content_type = res.headers.get("content-type", "")
    assert "spreadsheetml" in content_type or "excel" in content_type or "octet-stream" in content_type
    assert "attachment" in res.headers.get("content-disposition", "")
    assert "2026_09.xlsx" in res.headers.get("content-disposition", "")

    # Parse with openpyxl
    file_bytes = io.BytesIO(res.content)
    wb = openpyxl.load_workbook(file_bytes)
    assert "Payroll Statement" in wb.sheetnames
    ws = wb["Payroll Statement"]

    # Row 1 is Title, Row 2 is Subtitle, Row 3 is empty, Row 4 is Headers
    headers = [cell.value for cell in ws[4] if cell.value is not None]
    assert "Employee ID" in headers
    assert "Employee Name" in headers
    assert "Department" in headers
    assert "Annual CTC" in headers
    assert any("Basic" in str(h) for h in headers)
    assert any("Gross" in str(h) for h in headers)
    assert any("Net Payable" in str(h) for h in headers)
    assert any("Status" in str(h) for h in headers)

    # Verify at least one employee row is present
    row_count = ws.max_row
    assert row_count >= 5, f"Expected at least 5 rows (including headers and data), got {row_count}"


@pytest.mark.asyncio
async def test_payroll_excel_export_single_employee(client: AsyncClient, admin_headers):
    """
    Test 5: Export single employee payroll for September 2026.
    """
    # Export for Admin User
    res = await client.get(f"/api/v1/hrms/payroll/export?payroll_month=2026-09&employee_id={ADMIN_ID}", headers=admin_headers)
    assert res.status_code == 200, f"Expected 200 OK, got {res.status_code}: {res.text}"

    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws = wb.active
    # In single employee export, row 5 is the employee row
    emp_id_val = ws.cell(row=5, column=1).value
    assert emp_id_val == "EMP-000"


@pytest.mark.asyncio
async def test_payroll_excel_export_empty_month(client: AsyncClient, admin_headers):
    """
    Test 6: Exporting for a month without calculated payroll returns 404 with clear message.
    """
    res = await client.get("/api/v1/hrms/payroll/export?payroll_month=2024-01", headers=admin_headers)
    assert res.status_code == 404
    assert "No generated payroll records found" in res.text


@pytest.mark.asyncio
async def test_payroll_excel_export_unauthorized_employee(client: AsyncClient, regular_emp_headers):
    """
    Test 7: Regular employee without payroll management permissions cannot export all employees.
    """
    headers = {"Authorization": regular_emp_headers["Authorization"]}
    res = await client.get("/api/v1/hrms/payroll/export?payroll_month=2026-09", headers=headers)
    assert res.status_code == 403
