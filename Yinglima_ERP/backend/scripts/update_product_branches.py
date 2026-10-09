import asyncio
import json
import openpyxl
from sqlalchemy import text
from sqlalchemy.ext.asyncio import create_async_engine

EXCEL_PATH = r"D:\Om work1\ERP\Yinglima_ERP\doc\Master Planning Sheet China TO India (66).xlsx"
INHYMA_ORG_ID = "85c12bcb-10a1-4b1f-92ed-6ca72ac1b3b5"
AHMEDABAD_BRANCH_ID = "br_1786607390334_uitw"
INDORE_BRANCH_ID = "br_1786607543358_ugll"

DB_URLS = [
    ("mpvzjzunkiqchhhvxrza (backend/.env)", "postgresql+asyncpg://postgres.mpvzjzunkiqchhhvxrza:Inhyma%402026@aws-0-ap-south-1.pooler.supabase.com:6543/postgres?ssl=require"),
    ("ohjxbikcnlihiftodspe", "postgresql+asyncpg://postgres.ohjxbikcnlihiftodspe:s8idqqiYOX9oAZwk@aws-0-ap-south-1.pooler.supabase.com:6543/postgres?ssl=require"),
]

async def sync_branches():
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws_gj = wb["GJ Branch"]
    ws_mp = wb["MP Branch"]

    gj_keys = set()
    for r in range(2, ws_gj.max_row + 1):
        v = ws_gj.cell(r, 1).value
        if v and "Total" not in str(v):
            gj_keys.add(str(v).strip().lower())

    mp_keys = set()
    for r in range(2, ws_mp.max_row + 1):
        v = ws_mp.cell(r, 1).value
        if v and "Total" not in str(v):
            mp_keys.add(str(v).strip().lower())

    print(f"Loaded from Excel: {len(gj_keys)} GJ keys, {len(mp_keys)} MP keys.")

    connect_args = {"statement_cache_size": 0, "prepared_statement_name_func": lambda: ""}

    for db_label, db_url in DB_URLS:
        print(f"\n--- Syncing Product Branches on {db_label} ---")
        eng = create_async_engine(db_url, connect_args=connect_args)
        async with eng.connect() as conn:
            prods = (await conn.execute(text("SELECT id, product_name, product_name_tally FROM products"))).fetchall()
            print(f"Total products in DB: {len(prods)}")

            gj_count = 0
            mp_count = 0
            both_count = 0

            for pid, name, tally in prods:
                k1 = (name or "").strip().lower()
                k2 = (tally or "").strip().lower()

                in_gj = (k1 in gj_keys) or (k2 in gj_keys)
                in_mp = (k1 in mp_keys) or (k2 in mp_keys)

                b_ids = []
                if in_gj:
                    b_ids.append(AHMEDABAD_BRANCH_ID)
                if in_mp:
                    b_ids.append(INDORE_BRANCH_ID)

                # Fallback: if not found in either, assign both so it's not orphaned
                if not b_ids:
                    b_ids = [AHMEDABAD_BRANCH_ID, INDORE_BRANCH_ID]

                if len(b_ids) == 2:
                    both_count += 1
                elif b_ids[0] == AHMEDABAD_BRANCH_ID:
                    gj_count += 1
                else:
                    mp_count += 1

                org_json = json.dumps([INHYMA_ORG_ID])
                b_json = json.dumps(b_ids)

                await conn.execute(
                    text("UPDATE products SET organization_id = :org, organization_ids = CAST(:oids AS json), branch_ids = CAST(:bids AS json) WHERE id = :pid"),
                    {"org": INHYMA_ORG_ID, "oids": org_json, "bids": b_json, "pid": str(pid)}
                )

            await conn.commit()
            print(f"Successfully updated {len(prods)} products on {db_label}:")
            print(f"  * GJ Only (Ahmedabad): {gj_count}")
            print(f"  * MP Only (Indore): {mp_count}")
            print(f"  * Both Branches: {both_count}")

if __name__ == "__main__":
    asyncio.run(sync_branches())
