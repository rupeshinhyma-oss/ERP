import openpyxl

wb = openpyxl.load_workbook(r"D:\Om work1\ERP\Yinglima_ERP\doc\Master Planning Sheet China TO India (66).xlsx", data_only=True)

for name, total_row in (("GJ Branch", 243), ("MP Branch", 264)):
    ws = wb[name]
    rows = [(r, str(ws.cell(r, 1).value).strip()) for r in range(total_row + 1, ws.max_row + 1)
            if ws.cell(r, 1).value and str(ws.cell(r, 1).value).strip()]
    print(f"=== {name}: {len(rows)} spares, Excel rows {rows[0][0]} to {rows[-1][0]} ===")
    for r, v in rows[:8]:
        print(f"  Row {r}: {v}")
    print("  ...")
    for r, v in rows[-4:]:
        print(f"  Row {r}: {v}")
